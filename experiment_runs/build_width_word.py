import json, pathlib, collections, math
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
ROOT=pathlib.Path(__file__).resolve().parent
P=ROOT/'width_200/failure_analysis'
data=json.loads((P/'independent_comparison.json').read_text(encoding='utf-8'))
cases=data['cases']; byid={r['case_id']:r for r in cases}
OUT=ROOT.parent/'汇报材料/多智能体宽度200条审核汇总.docx'
OUT.parent.mkdir(exist_ok=True)
doc=Document(); section=doc.sections[0]
section.page_width=Cm(21);section.page_height=Cm(29.7)
section.top_margin=Cm(1.8);section.bottom_margin=Cm(1.7);section.left_margin=Cm(1.8);section.right_margin=Cm(1.8)
section.header_distance=Cm(.7);section.footer_distance=Cm(.7)
for grid in section._sectPr.findall(qn('w:docGrid')):section._sectPr.remove(grid)
for name,size in [('Normal',10.5),('Title',23),('Heading 1',16),('Heading 2',12)]:
    style=doc.styles[name];style.font.name='Microsoft YaHei';style.font.size=Pt(size);style.font.color.rgb=RGBColor(0,0,0)
    style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
    style.paragraph_format.space_after=Pt(6)
    for border in style.element.xpath('./w:pPr/w:pBdr'):border.getparent().remove(border)
    snap=OxmlElement('w:snapToGrid');snap.set(qn('w:val'),'0');style.element.get_or_add_pPr().append(snap)
doc.styles['Normal'].paragraph_format.line_spacing=Pt(14)
doc.styles['Normal'].paragraph_format.space_after=Pt(5)
for name in ('Title','Heading 1','Heading 2'):doc.styles[name].paragraph_format.line_spacing=1.05
header=section.header.paragraphs[0];header.text='实验结果核对报告   多智能体宽度组'
header.runs[0].font.size=Pt(8);header.runs[0].font.color.rgb=RGBColor.from_string('64748B')
footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('统计日期 2026年10月19日    第 ')
f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');footer._p.append(f);footer.add_run(' 页')
for r in footer.runs:r.font.size=Pt(8)
def p(text='',bold=False,size=None):
    x=doc.add_paragraph();r=x.add_run(text);r.bold=bold
    if size:r.font.size=Pt(size)
    return x
def h(text,level=1):doc.add_heading(text,level)
def table(headers,rows,widths,size=9.5):
    t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
    for c,w in zip(t.columns,widths):c.width=Cm(w)
    for c,txt in zip(t.rows[0].cells,headers):c.text=txt
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        cells=t.add_row().cells
        for c,txt in zip(cells,row):c.text=str(txt)
    for i,row in enumerate(t.rows):
        cant=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(cant)
        for c,w in zip(row.cells,widths):
            c.width=Cm(w);pr=c._tc.get_or_add_tcPr()
            margins=OxmlElement('w:tcMar')
            for side in ('top','left','bottom','right'):
                el=OxmlElement('w:'+side);el.set(qn('w:w'),'40' if side in ('top','bottom') else '65');el.set(qn('w:type'),'dxa');margins.append(el)
            pr.append(margins)
            if i==0:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EDF2F7');pr.append(sh)
            for para in c.paragraphs:
                para.paragraph_format.space_after=Pt(1);para.paragraph_format.space_before=Pt(0);para.paragraph_format.line_spacing=Pt(size+2)
                for r in para.runs:r.font.size=Pt(size);r.bold=i==0
    borders=OxmlElement('w:tblBorders')
    for side in ('top','left','bottom','right','insideH','insideV'):
        e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D7DFE7');borders.append(e)
    t._tbl.tblPr.append(borders)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)
    return t
def page():doc.add_page_break()
def ids(items):return '、'.join(i.removeprefix('req_') for i in items)
codes={'需求车道在真实道路中不存在':'A','目标车道存在但DSL选错':'B','选道正确但DSL自检拒绝（需复核其余字段）':'C','原始地图文件缺失':'D','现有检查未定位原因，需人工复核':'E'}

doc.add_heading('多智能体宽度200条审核汇总',0)
p('逐条比较 DSL 和工具调用与数据库标准答案',True,12)
p('本报告覆盖 width_200_trace.jsonl 中的 200 条宽度任务，保留人工通过与未通过记录，并直接读取每条 trace 的最终 DSL、工具调用及数据库候选答案进行独立比较。统计不采用前端正确率卡片，也未修改原始结果或标准答案。')
table(['统计项目','通过或一致','不通过或不一致','其他','占全部200条比例'],[
['人工审核','108','92','0','通过率 54.0%'],
['DSL 与标准答案比较','121','79','0','一致率 60.5%'],
['工具调用与标准答案比较','111','83','6 条无基准','一致率 55.5%'],
['任务执行状态','162 done','38 error','0','完成率 81.0%'],
],[4.3,2.4,3.3,2.8,4.6])
p('DSL 的 79 条未通过包含 78 条字段不一致和 1 条无输出。工具调用的 83 条未通过包含 45 条字段不一致和 38 条无输出；其余 6 条因缺少已批准的标准工具答案，不能判定一致或不一致。',size=10)
h('统计口径',2)
p('主表分母统一为本组全部 200 条任务。工具调用若排除 6 条无基准任务，已判定 194 条中的一致率为 111÷194＝57.2%。人工审核、标准答案一致性与真实道路修改是否正确，是三个不同指标。')
h('逐条比较方法',2)
p('每条任务合并 trace 中同一步骤的增量记录，读取最后保留的宽度 DSL 与实际 tool_call_json。本组每条有输出的任务均只有一个宽度步骤。未执行工具调用的任务，不用计划或旧 DSL 参考命令代替实际调用。')
p('标准答案优先选取该任务真实道路对应的 width 候选；没有真实道路候选时，使用原需求 Road 21 的候选，并根据该条已有映射统一道路编号。共 94 条采用真实道路候选，106 条采用需求道路候选。车道编号和宽度常量不会随道路编号自动改写。')
p('DSL 按字段比较；忽略 JSON 对象键顺序。工具调用独立比较同一任务相关的所有 approved 候选，不要求先唯一关联某一份 DSL。公式按表达式结构比较，忽略空格、多余括号及 2 与 2.0 这类数值写法；不自动合并代数等价公式。调用顺序、车道、公式常量及 s、ms、sh 等参数保持核对。',size=10)

page();h('DSL 和工具调用的独立核对结果')
table(['阶段','结果分类','条数','占200条比例'],[
['DSL','与相关标准候选一致',121,'60.5%'],['DSL','lane_ids 字段不同',56,'28.0%'],['DSL','operations 数量不同',22,'11.0%'],['DSL','无 DSL 输出',1,'0.5%'],
['工具调用','与相关 approved 候选一致',111,'55.5%'],['工具调用','字段不同',45,'22.5%'],['工具调用','无实际工具调用',38,'19.0%'],['工具调用','没有可用 approved 基准',6,'3.0%'],
],[2.2,9,2,4.2])
p('45 条工具字段不一致中，44 条存在 li 车道差异，43 条存在 formula 差异，两者重叠 42 条，不能相加为失败总数。未发现这 45 条的 s、ms、sh 等其他字段差异。')
h('需要单独说明的记录',2)
p('req_000012：工具公式为 2.0，标准为 2，数值和表达式含义相同，本报告计为一致；如果仅比较去空格后的公式原文，工具一致数会是 110 条。')
p('req_000019：DSL 匹配，但工具公式是 3.3−3.7896，标准是 3.3−3.4903。后一个常量来自旧道路宽度，前一个来自本次真实道路上下文，因此这条是“与现存基准不一致”，不能只据此断言新公式计算错误；该条还存在选道语义问题。')
p('无工具基准的 6 条：'+ids([r['case_id'] for r in cases if r['tool']['decision']=='no_baseline'])+'。这些任务均有工具输出，但在本次所选 DSL 候选下没有 approved 工具答案，报告不将其直接计为工具错误。')
h('标准答案也需要复核',2)
p('req_000011、req_000014、req_000040：DSL 和工具调用均与当前标准答案一致，人工记录却为未通过。按项目现有选道规则，这三条仍可能选错右侧车道。这说明“匹配标准答案”只代表一致性，不能证明需求执行正确。')
p('人工通过的 108 条在本次独立比较中均为 DSL 一致且工具调用一致。人工未通过的 92 条中，13 条 DSL 仍一致，3 条工具调用仍一致；其余情况见逐条附录。')
h('执行结果与比较结果的边界',2)
p('本报告的工具一致性比较 tool_call_json 的名称、参数和公式，不把 resolved_result、内部命令或道路几何结果当作工具调用标准答案。请求目标超出宽度限值、联动道路范围及实际文件中的最终宽度，仍需结合执行后地图另审。done 只表示流程结束，不能替代正确性结论。',size=10)

page();h('人工未通过原因汇总')
p('人工未通过 92 条的备注均为空。下表为根据真实道路上下文、最终 DSL、自检记录和执行错误做的独立原因筛查，非用户已填写的拒绝理由。类别互斥；同一任务还可能同时存在公式或最终宽度问题。',size=10)
table(['代码','筛查原因','条数','占92条比例'],[
['A','需求所指车道在映射后的真实道路中不存在',44,'47.8%'],['B','按项目选道规则应有目标车道，但最终 DSL 选错',34,'37.0%'],['C','选道符合项目规则，但 DSL 自检拒绝，需复核其余字段',9,'9.8%'],['D','原始地图文件缺失，未进入正常修改',1,'1.1%'],['E','现有筛查未定位充分原因，需要人工复核',4,'4.3%'],
],[1.2,11.1,1.8,3.3])
p('A 和 B 合计 78 条，占未通过记录的 84.8%，集中在需求车道与真实道路的兼容性及车道选择。A 类中 22 条没有目标车道却改选其他车道；另外 22 条没有生成对应操作。左右规则依据本项目提示词和该条 lane_ids 顺序，不能套用画面左右或通用编号猜测。',size=10)
p('A 类例子 req_000004：需求修改 lane 2，真实 road_3 仅有 [1,−1]，DSL 改为 lane_1。C 类例子 req_000095：最终自检文本为 OK，但 passed=false，执行错误为 Width DSL check failed: OK。D 类为 req_000460，缺少原始 XODR 文件。',size=10)
p('E 类为 req_000088、req_000439、req_000471、req_000482。当前筛查发现选道和主要数值关系基本符合，尚不能确定原人工拒绝原因；其中 439、471 有多条道路联动，需要核对范围，不能凭联动条数直接判错。',size=10)
h('人工通过任务清单',2)
passed=[r['case_id'] for r in cases if r['human']=='pass']
p('以下编号均省略 req_ 前缀，共 108 条；在本次独立核对中两项均一致。',size=9)
for i in range(0,len(passed),12):p(ids(passed[i:i+12]),size=8.4)
p('数据来源：request_tests.sqlite3 的 case_reviews、DSL 候选及工具答案；width_200_trace.jsonl；各 cases 目录的 input.json、execution_result.json 与道路上下文。运行标识 width_200_real_20261017；模型 deepseek-chat；agent_flow 为 normal。',size=8)

def status(r,s):
    return {'pass':'一致','fail':'不一致','no_output':'无输出','no_baseline':'无基准'}[r[s]['decision']]
def desc(r):
    ds=r['dsl'];ts=r['tool'];text=[]
    if ds['decision']=='fail':
        lane=next((d for d in ds['diffs'] if 'lane_ids[' in d['field']),None)
        if lane:text.append('DSL '+str(lane['actual']).replace('lane_','')+'≠标准'+str(lane['standard']).replace('lane_',''))
        else:text.append('DSL 操作数不同')
    elif ds['decision']=='no_output':text.append('缺原始地图')
    if ts['decision']=='fail':
        fields=[d['field'] for d in ts['diffs']]
        text.append('工具'+('车道' if any('.li[' in x for x in fields) else '')+('及公式' if any('.li[' in x for x in fields) and any('.formula' in x for x in fields) else '公式' if any('.formula' in x for x in fields) else '')+'不同')
    elif ts['decision']=='no_output':text.append('未调用工具')
    elif ts['decision']=='no_baseline':text.append('缺 approved 工具答案')
    if not text:text.append('匹配答案 仍需审选道')
    return '；'.join(text)
failed=[r for r in cases if r['human']=='fail']
for offset in range(0,len(failed),23):
    page();h(f'人工未通过逐条明细 {offset//23+1}')
    p('本页全部为人工未通过。DSL 与工具列来自独立字段比较；A 至 E 对应上一节原因分类。编号保留完整形式，便于搜索审核器。',size=9)
    rows=[]
    for r in failed[offset:offset+23]:
        rows.append([r['case_id'],r['road'],status(r,'dsl'),status(r,'tool'),codes[r['triage']['category']],desc(r)])
    table(['任务编号','真实路','DSL','工具','原因','核对差异或待核事项'],rows,[3.0,1.2,1.5,1.5,1.1,9.1],9)
    p('“标准”指本次选择的数据库候选，不表示其已完成真实道路语义校验。原始数据库内可存在多个候选，任一相关候选匹配即计为一致；真实道路候选优先。',size=9)
doc.save(OUT)
print(OUT)
