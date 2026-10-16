import argparse
import json
import os
import re
import sqlite3
import tempfile
from pathlib import Path

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_WORKBOOK_PATH = PROJECT_ROOT / "testfucs" / "run_data" / "请求测试.xlsx"
DEFAULT_DATABASE_PATH = (
    PROJECT_ROOT / "testfucs" / "agent_trace_viewer" / "request_tests.sqlite3"
)
SCHEMA_VERSION = 2
AGENT_TYPES = ("width", "slope", "curve")
ANSWER_COLUMNS = {
    "width": "width answer",
    "slope": "slope answer",
    "curve": "curve answer",
}
REQUIRED_COLUMNS = (
    "输入",
    "算子",
    "类型",
    "width",
    "slope",
    "curve",
    "width answer",
    "slope answer",
    "curve answer",
)


SCHEMA_SQL = """
CREATE TABLE test_cases (
    case_id TEXT PRIMARY KEY,
    source_row INTEGER NOT NULL UNIQUE,
    input_text TEXT NOT NULL UNIQUE,
    operator_text TEXT,
    operator_types_json TEXT NOT NULL CHECK (json_valid(operator_types_json)),
    case_type TEXT,
    road_context_json TEXT NOT NULL CHECK (json_valid(road_context_json)),
    created_at TEXT,
    updated_at TEXT
);

CREATE TABLE dsl_answers (
    dsl_answer_id INTEGER PRIMARY KEY,
    answer_set_id TEXT NOT NULL,
    case_id TEXT NOT NULL,
    subtask_key TEXT NOT NULL,
    agent_type TEXT NOT NULL CHECK (agent_type IN ('width', 'slope', 'curve')),
    road_id TEXT,
    subtask_description TEXT,
    dsl_json TEXT NOT NULL CHECK (json_valid(dsl_json)),
    review_status TEXT CHECK (
        review_status IS NULL
        OR review_status IN ('draft', 'approved', 'rejected', 'stale')
    ),
    metadata_json TEXT NOT NULL CHECK (json_valid(metadata_json)),
    created_at TEXT,
    updated_at TEXT,
    FOREIGN KEY (case_id) REFERENCES test_cases(case_id) ON DELETE CASCADE,
    UNIQUE (answer_set_id, subtask_key)
);

CREATE TABLE tool_call_answers (
    tool_answer_id INTEGER PRIMARY KEY,
    dsl_answer_id INTEGER NOT NULL,
    variant_index INTEGER NOT NULL CHECK (variant_index > 0),
    tool_call_json TEXT NOT NULL CHECK (json_valid(tool_call_json)),
    resolved_result_json TEXT CHECK (
        resolved_result_json IS NULL OR json_valid(resolved_result_json)
    ),
    internal_command_json TEXT CHECK (
        internal_command_json IS NULL OR json_valid(internal_command_json)
    ),
    random_seed INTEGER,
    review_status TEXT CHECK (
        review_status IS NULL
        OR review_status IN ('draft', 'approved', 'rejected', 'stale')
    ),
    metadata_json TEXT NOT NULL CHECK (json_valid(metadata_json)),
    created_at TEXT,
    updated_at TEXT,
    FOREIGN KEY (dsl_answer_id) REFERENCES dsl_answers(dsl_answer_id)
        ON DELETE CASCADE,
    UNIQUE (dsl_answer_id, variant_index)
);

CREATE TABLE case_reviews (
    review_id INTEGER PRIMARY KEY,
    case_id TEXT NOT NULL,
    run_id TEXT NOT NULL,
    task_id TEXT NOT NULL,
    decision TEXT NOT NULL CHECK (
        decision IN ('pending', 'pass', 'question', 'fail')
    ),
    note TEXT NOT NULL DEFAULT '',
    review_source TEXT NOT NULL CHECK (review_source IN ('human', 'auto')),
    trace_source TEXT,
    metadata_json TEXT NOT NULL CHECK (json_valid(metadata_json)),
    created_at TEXT,
    updated_at TEXT,
    FOREIGN KEY (case_id) REFERENCES test_cases(case_id) ON DELETE CASCADE,
    UNIQUE (case_id, run_id, task_id)
);

CREATE INDEX idx_dsl_answers_case
    ON dsl_answers(case_id, answer_set_id, agent_type);
CREATE INDEX idx_tool_call_answers_dsl
    ON tool_call_answers(dsl_answer_id);
CREATE INDEX idx_case_reviews_run
    ON case_reviews(run_id, task_id);
"""


CASE_REVIEWS_SCHEMA_SQL = """
CREATE TABLE case_reviews (
    review_id INTEGER PRIMARY KEY,
    case_id TEXT NOT NULL,
    run_id TEXT NOT NULL,
    task_id TEXT NOT NULL,
    decision TEXT NOT NULL CHECK (
        decision IN ('pending', 'pass', 'question', 'fail')
    ),
    note TEXT NOT NULL DEFAULT '',
    review_source TEXT NOT NULL CHECK (review_source IN ('human', 'auto')),
    trace_source TEXT,
    metadata_json TEXT NOT NULL CHECK (json_valid(metadata_json)),
    created_at TEXT,
    updated_at TEXT,
    FOREIGN KEY (case_id) REFERENCES test_cases(case_id) ON DELETE CASCADE,
    UNIQUE (case_id, run_id, task_id)
);
CREATE INDEX idx_case_reviews_run
    ON case_reviews(run_id, task_id);
"""


def connect_database(database_path):
    connection = sqlite3.connect(Path(database_path))
    connection.row_factory = sqlite3.Row
    connection.execute("PRAGMA foreign_keys = ON")
    connection.execute("PRAGMA busy_timeout = 5000")
    migrate_database(connection)
    return connection


def migrate_database(connection):
    version = connection.execute("PRAGMA user_version").fetchone()[0]
    application_tables = connection.execute(
        "SELECT name FROM sqlite_master WHERE type = 'table' AND name = 'test_cases'"
    ).fetchall()
    if not application_tables:
        return
    if version == 1:
        with connection:
            connection.executescript(CASE_REVIEWS_SCHEMA_SQL)
            connection.execute(f"PRAGMA user_version = {SCHEMA_VERSION}")
        return
    if version != SCHEMA_VERSION:
        raise RuntimeError(
            f"Unsupported database schema version: {version}; expected {SCHEMA_VERSION}."
        )


def initialize_database(connection):
    existing = connection.execute(
        "SELECT name FROM sqlite_master WHERE type = 'table' AND name NOT LIKE 'sqlite_%'"
    ).fetchall()
    if existing:
        raise ValueError("Database already contains application tables.")
    connection.executescript(SCHEMA_SQL)
    connection.execute(f"PRAGMA user_version = {SCHEMA_VERSION}")


def _json_text(value):
    return json.dumps(value, ensure_ascii=False, separators=(",", ":"))


def _parse_operator_types(value):
    if value is None or not str(value).strip():
        return []
    text = str(value).strip().lower()
    if text == "all":
        return list(AGENT_TYPES)
    tokens = [
        token.strip()
        for token in re.split(r"[、,，/\s]+", text)
        if token.strip()
    ]
    unknown = [token for token in tokens if token not in AGENT_TYPES]
    if unknown:
        raise ValueError(f"Unsupported operator value: {value}")
    return [agent_type for agent_type in AGENT_TYPES if agent_type in tokens]


def _parse_answer_entries(cell_value, source_row, answer_column):
    if cell_value is None or not str(cell_value).strip():
        return []
    try:
        parsed = json.loads(cell_value) if isinstance(cell_value, str) else cell_value
    except json.JSONDecodeError as exc:
        raise ValueError(
            f"Invalid JSON at Sheet1 row {source_row}, column {answer_column}."
        ) from exc

    if isinstance(parsed, dict):
        records = [parsed]
    elif isinstance(parsed, list):
        if all(isinstance(item, dict) and "operation" in item for item in parsed):
            records = [{"operations": parsed}]
        else:
            records = parsed
    else:
        raise ValueError(
            f"Unsupported answer at Sheet1 row {source_row}, column {answer_column}."
        )

    entries = []
    for record_index, record in enumerate(records, start=1):
        if not isinstance(record, dict):
            raise ValueError(
                f"Answer entry {record_index} at Sheet1 row {source_row}, "
                f"column {answer_column} must be an object."
            )
        if "value" in record:
            payload = record["value"]
            metadata = {key: value for key, value in record.items() if key != "value"}
        else:
            payload = record
            metadata = {}
        if isinstance(payload, str):
            try:
                payload = json.loads(payload)
            except json.JSONDecodeError as exc:
                raise ValueError(
                    f"Invalid nested JSON at Sheet1 row {source_row}, "
                    f"column {answer_column}, entry {record_index}."
                ) from exc
        if not isinstance(payload, dict):
            raise ValueError(
                f"DSL payload at Sheet1 row {source_row}, column {answer_column}, "
                f"entry {record_index} must be an object."
            )
        operations = payload.get("operations")
        if not isinstance(operations, list):
            raise ValueError(
                f"DSL payload at Sheet1 row {source_row}, column {answer_column}, "
                f"entry {record_index} must contain an operations array."
            )
        entries.append((payload, metadata))
    return entries


def _group_dsl_by_road(payload):
    operations = payload["operations"]
    if not operations:
        return [(None, payload)]

    grouped = {}
    order = []
    for operation in operations:
        if not isinstance(operation, dict):
            road_id = None
        else:
            road_id = operation.get("road_id")
            road_id = None if road_id is None else str(road_id)
        group_key = road_id or "__none__"
        if group_key not in grouped:
            grouped[group_key] = []
            order.append(group_key)
        grouped[group_key].append(operation)

    result = []
    for group_key in order:
        grouped_payload = dict(payload)
        grouped_payload["operations"] = grouped[group_key]
        result.append((None if group_key == "__none__" else group_key, grouped_payload))
    return result


def _insert_case(connection, case_id, source_row, row_data):
    input_text = str(row_data["输入"]).strip()
    operator_text = row_data.get("算子")
    operator_types = _parse_operator_types(operator_text)
    road_context = {
        "width": row_data.get("width"),
        "slope": row_data.get("slope"),
        "curve": row_data.get("curve"),
    }
    connection.execute(
        """
        INSERT INTO test_cases (
            case_id, source_row, input_text, operator_text, operator_types_json,
            case_type, road_context_json, created_at, updated_at
        ) VALUES (?, ?, ?, ?, ?, ?, ?, NULL, NULL)
        """,
        (
            case_id,
            source_row,
            input_text,
            operator_text,
            _json_text(operator_types),
            row_data.get("类型"),
            _json_text(road_context),
        ),
    )


def _build_dsl_answer_records(case_id, source_row, row_data):
    records = []
    source_cells = 0
    for agent_type, answer_column in ANSWER_COLUMNS.items():
        entries = _parse_answer_entries(
            row_data.get(answer_column), source_row, answer_column
        )
        if entries:
            source_cells += 1
        for candidate_index, (payload, metadata) in enumerate(entries, start=1):
            answer_set_id = f"{case_id}_set_{candidate_index}"
            for subtask_index, (road_id, grouped_payload) in enumerate(
                _group_dsl_by_road(payload), start=1
            ):
                records.append(
                    {
                        "answer_set_id": answer_set_id,
                        "case_id": case_id,
                        "subtask_key": f"{agent_type}_{subtask_index}",
                        "agent_type": agent_type,
                        "road_id": road_id,
                        "subtask_description": None,
                        "dsl_json": _json_text(grouped_payload),
                        "review_status": None,
                        "metadata_json": _json_text(metadata),
                        "created_at": metadata.get("created_at"),
                        "updated_at": metadata.get("updated_at"),
                    }
                )
    return records, source_cells


def _insert_dsl_answers(connection, case_id, source_row, row_data):
    records, source_cells = _build_dsl_answer_records(
        case_id, source_row, row_data
    )
    connection.executemany(
        """
        INSERT INTO dsl_answers (
            answer_set_id, case_id, subtask_key, agent_type, road_id,
            subtask_description, dsl_json, review_status, metadata_json,
            created_at, updated_at
        ) VALUES (
            :answer_set_id, :case_id, :subtask_key, :agent_type, :road_id,
            NULL, :dsl_json, NULL, :metadata_json, :created_at, :updated_at
        )
        """,
        records,
    )
    return len(records), source_cells


def migrate_sheet1(workbook_path, database_path, overwrite=False):
    workbook_path = Path(workbook_path).resolve()
    database_path = Path(database_path).resolve()
    if not workbook_path.is_file():
        raise FileNotFoundError(f"Workbook not found: {workbook_path}")
    if database_path.exists() and not overwrite:
        raise FileExistsError(f"Database already exists: {database_path}")
    database_path.parent.mkdir(parents=True, exist_ok=True)

    file_descriptor, temp_name = tempfile.mkstemp(
        prefix=f"{database_path.name}.", suffix=".tmp", dir=database_path.parent
    )
    os.close(file_descriptor)
    temp_path = Path(temp_name)
    connection = None
    workbook = None
    try:
        connection = connect_database(temp_path)
        initialize_database(connection)
        workbook = load_workbook(workbook_path, read_only=True, data_only=True)
        if "Sheet1" not in workbook.sheetnames:
            raise ValueError("Workbook must contain Sheet1.")
        worksheet = workbook["Sheet1"]
        header_row = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True))
        headers = {value: index for index, value in enumerate(header_row) if value}
        missing = [column for column in REQUIRED_COLUMNS if column not in headers]
        if missing:
            raise ValueError(f"Sheet1 is missing columns: {', '.join(missing)}")

        case_count = 0
        dsl_answer_count = 0
        answer_source_cell_count = 0
        with connection:
            for source_row, row in enumerate(
                worksheet.iter_rows(min_row=2, values_only=True), start=2
            ):
                row_data = {
                    header: row[index] if index < len(row) else None
                    for header, index in headers.items()
                }
                if row_data.get("输入") is None or not str(row_data["输入"]).strip():
                    continue
                case_count += 1
                case_id = f"req_{case_count:06d}"
                _insert_case(connection, case_id, source_row, row_data)
                inserted, source_cells = _insert_dsl_answers(
                    connection, case_id, source_row, row_data
                )
                dsl_answer_count += inserted
                answer_source_cell_count += source_cells

        verification = verify_database(connection)
        if verification["test_cases"] != case_count:
            raise RuntimeError("Migrated test case count does not match Sheet1.")
        if verification["dsl_answers"] != dsl_answer_count:
            raise RuntimeError("Migrated DSL answer count does not match inserts.")
        if verification["tool_call_answers"] != 0:
            raise RuntimeError("Tool answers must remain empty during Sheet1 migration.")

        connection.close()
        connection = None
        workbook.close()
        workbook = None
        os.replace(temp_path, database_path)
        return {
            "database": str(database_path),
            "sheet": "Sheet1",
            "test_cases": case_count,
            "answer_source_cells": answer_source_cell_count,
            "dsl_answers": dsl_answer_count,
            "tool_call_answers": 0,
        }
    finally:
        if workbook is not None:
            workbook.close()
        if connection is not None:
            connection.close()
        if temp_path.exists():
            temp_path.unlink()


def validate_resolved_result(agent_type, resolved_result):
    if resolved_result is None:
        return
    if agent_type in {"width", "slope"}:
        if not isinstance(resolved_result, list):
            raise ValueError(f"{agent_type} resolved result must be a call array.")
        return
    if agent_type != "curve":
        raise ValueError(f"Unsupported agent type: {agent_type}")
    if not isinstance(resolved_result, dict):
        raise ValueError("Curve resolved result must be an object.")
    control_points = resolved_result.get("optimized_control_points")
    if not isinstance(control_points, list) or len(control_points) != 4:
        raise ValueError("Curve resolved result must contain four control points.")
    for point in control_points:
        if (
            not isinstance(point, list)
            or len(point) != 2
            or any(not isinstance(value, (int, float)) for value in point)
        ):
            raise ValueError("Each curve control point must be a numeric [x, y] pair.")


def insert_tool_call_answer(
    connection,
    dsl_answer_id,
    variant_index,
    tool_calls,
    resolved_result=None,
    internal_commands=None,
    random_seed=None,
    review_status=None,
    metadata=None,
):
    dsl_answer = connection.execute(
        "SELECT agent_type FROM dsl_answers WHERE dsl_answer_id = ?",
        (dsl_answer_id,),
    ).fetchone()
    if dsl_answer is None:
        raise ValueError(f"DSL answer does not exist: {dsl_answer_id}")
    if not isinstance(tool_calls, list):
        raise ValueError("Tool calls must be an array.")
    if internal_commands is not None and (
        not isinstance(internal_commands, list)
        or any(not isinstance(command, str) for command in internal_commands)
    ):
        raise ValueError("Internal commands must be a string array.")
    validate_resolved_result(dsl_answer["agent_type"], resolved_result)
    cursor = connection.execute(
        """
        INSERT INTO tool_call_answers (
            dsl_answer_id, variant_index, tool_call_json, resolved_result_json,
            internal_command_json, random_seed, review_status, metadata_json,
            created_at, updated_at
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?,
                  strftime('%Y-%m-%dT%H:%M:%fZ', 'now'),
                  strftime('%Y-%m-%dT%H:%M:%fZ', 'now'))
        """,
        (
            dsl_answer_id,
            variant_index,
            _json_text(tool_calls),
            None if resolved_result is None else _json_text(resolved_result),
            None if internal_commands is None else _json_text(internal_commands),
            random_seed,
            review_status,
            _json_text(metadata or {}),
        ),
    )
    return cursor.lastrowid


def _normalize_dsl_road_id(value):
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None
    if text.lower().startswith("road_"):
        return f"road_{text[5:]}"
    return f"road_{text}"


def _matched_alignment_road_id(alignment):
    if not alignment:
        return None
    status = str(alignment.get("match_status") or "").strip().lower()
    if status != "matched":
        return None
    return _normalize_dsl_road_id(alignment.get("road_id"))


def _synchronize_dsl_road_ids(value, road_id):
    if not road_id:
        return value
    if isinstance(value, list):
        return [_synchronize_dsl_road_ids(item, road_id) for item in value]
    if not isinstance(value, dict):
        return value
    synchronized = {}
    for key, child in value.items():
        if key == "road_id" and child is not None and str(child).strip():
            synchronized[key] = road_id
        else:
            synchronized[key] = _synchronize_dsl_road_ids(child, road_id)
    return synchronized


def _synchronize_dsl_entries(entries, road_id):
    if not road_id:
        return entries
    synchronized = []
    for entry in entries:
        item = dict(entry)
        if "value" in item:
            item["value"] = _synchronize_dsl_road_ids(item["value"], road_id)
        else:
            item = _synchronize_dsl_road_ids(item, road_id)
        synchronized.append(item)
    return synchronized


def _synchronize_tool_road_ids(value, road_id):
    if not road_id:
        return value
    tool_road_id = road_id[5:] if road_id.lower().startswith("road_") else road_id
    if isinstance(value, list):
        return [_synchronize_tool_road_ids(item, road_id) for item in value]
    if not isinstance(value, dict):
        return value
    synchronized = {}
    for key, child in value.items():
        if (
            key in {"road_id", "id"}
            and child is not None
            and str(child).strip()
        ):
            child_text = str(child).strip().lower()
            synchronized[key] = (
                road_id if child_text.startswith("road_") else tool_road_id
            )
        else:
            synchronized[key] = _synchronize_tool_road_ids(child, road_id)
    return synchronized


def _synchronize_internal_command_road_ids(value, road_id):
    if not road_id:
        return value
    tool_road_id = road_id[5:] if road_id.lower().startswith("road_") else road_id
    if isinstance(value, list):
        return [
            _synchronize_internal_command_road_ids(item, road_id)
            for item in value
        ]
    if not isinstance(value, str):
        return value
    return re.sub(
        r"(^|\n)(\s*(?:width|slope|curve)\s+id=)[^\s]+",
        rf"\g<1>\g<2>{tool_road_id}",
        value,
        flags=re.IGNORECASE,
    )


def _synchronize_tool_answer(answer, road_id):
    if not road_id:
        return answer
    synchronized = dict(answer)
    for key in ("tool_calls", "resolved_result"):
        if key in synchronized:
            synchronized[key] = _synchronize_tool_road_ids(
                synchronized[key], road_id
            )
    if "internal_commands" in synchronized:
        synchronized["internal_commands"] = _synchronize_internal_command_road_ids(
            synchronized["internal_commands"], road_id
        )
    return synchronized


def list_dsl_answer_subtasks(connection, case_id, agent_type):
    if agent_type not in AGENT_TYPES:
        raise ValueError(f"Unsupported agent type: {agent_type}")
    _find_case(connection, case_id=case_id)
    alignment = list_case_road_alignments(connection, [case_id]).get(case_id)
    aligned_road_id = _matched_alignment_road_id(alignment)
    rows = connection.execute(
        """
        SELECT *
        FROM dsl_answers
        WHERE case_id = ? AND agent_type = ?
        ORDER BY dsl_answer_id
        """,
        (case_id, agent_type),
    ).fetchall()
    subtasks = []
    for row in rows:
        subtasks.append(
            {
                "dsl_answer_id": row["dsl_answer_id"],
                "answer_set_id": row["answer_set_id"],
                "subtask_key": row["subtask_key"],
                "agent_type": row["agent_type"],
                "road_id": aligned_road_id or row["road_id"],
                "dsl": _synchronize_dsl_road_ids(
                    _load_json(row["dsl_json"], {}), aligned_road_id
                ),
                "review_status": row["review_status"],
                "metadata": _load_json(row["metadata_json"], {}),
                "tool_answers": list_tool_call_answers(
                    connection, row["dsl_answer_id"], road_id=aligned_road_id
                ),
            }
        )
    return subtasks


def list_case_road_alignments(connection, case_ids=None):
    """Return the database-owned actual road mapping for each request case."""
    table_exists = connection.execute(
        """
        SELECT 1
        FROM sqlite_master
        WHERE type = 'table' AND name = 'test_case_road_alignments'
        """
    ).fetchone()
    if table_exists is None:
        return {}

    normalized_case_ids = None
    if case_ids is not None:
        if not isinstance(case_ids, (list, tuple, set)):
            raise ValueError("Road alignment case_ids must be an array.")
        normalized_case_ids = []
        seen = set()
        for case_id in case_ids:
            value = str(case_id or "").strip()
            if value and value not in seen:
                normalized_case_ids.append(value)
                seen.add(value)
        if not normalized_case_ids:
            return {}

    rows = []
    if normalized_case_ids is None:
        rows = connection.execute(
            """
            SELECT case_id, match_status, scene_file_name, map_file_name, road_id
            FROM test_case_road_alignments
            ORDER BY case_id
            """
        ).fetchall()
    else:
        for offset in range(0, len(normalized_case_ids), 400):
            chunk = normalized_case_ids[offset : offset + 400]
            placeholders = ",".join("?" for _ in chunk)
            rows.extend(
                connection.execute(
                    f"""
                    SELECT case_id, match_status, scene_file_name, map_file_name, road_id
                    FROM test_case_road_alignments
                    WHERE case_id IN ({placeholders})
                    ORDER BY case_id
                    """,
                    chunk,
                ).fetchall()
            )

    return {
        row["case_id"]: {
            "case_id": row["case_id"],
            "match_status": row["match_status"],
            "scene_file_name": row["scene_file_name"],
            "map_file_name": row["map_file_name"],
            "road_id": row["road_id"],
        }
        for row in rows
    }


def list_auto_review_baselines(connection, case_ids):
    if not isinstance(case_ids, list):
        raise ValueError("Auto-review case_ids must be an array.")
    normalized_case_ids = []
    seen = set()
    for case_id in case_ids:
        value = str(case_id or "").strip()
        if not value:
            raise ValueError("Auto-review case_id must not be empty.")
        if value not in seen:
            normalized_case_ids.append(value)
            seen.add(value)
    if len(normalized_case_ids) > 2000:
        raise ValueError("Auto-review batch is too large.")

    road_alignments = list_case_road_alignments(connection, normalized_case_ids)
    result = {
        case_id: {
            **{agent_type: [] for agent_type in AGENT_TYPES},
            "road_alignment": road_alignments.get(case_id),
        }
        for case_id in normalized_case_ids
    }
    dsl_rows = []
    for offset in range(0, len(normalized_case_ids), 400):
        chunk = normalized_case_ids[offset : offset + 400]
        placeholders = ",".join("?" for _ in chunk)
        dsl_rows.extend(
            connection.execute(
                f"""
                SELECT * FROM dsl_answers
                WHERE case_id IN ({placeholders})
                ORDER BY dsl_answer_id
                """,
                chunk,
            ).fetchall()
        )

    tool_answers = {}
    dsl_answer_ids = [row["dsl_answer_id"] for row in dsl_rows]
    for offset in range(0, len(dsl_answer_ids), 400):
        chunk = dsl_answer_ids[offset : offset + 400]
        placeholders = ",".join("?" for _ in chunk)
        rows = connection.execute(
            f"""
            SELECT * FROM tool_call_answers
            WHERE dsl_answer_id IN ({placeholders})
            ORDER BY dsl_answer_id, variant_index
            """,
            chunk,
        ).fetchall()
        for row in rows:
            tool_answers.setdefault(row["dsl_answer_id"], []).append(
                _tool_answer_dict(row)
            )

    for row in dsl_rows:
        aligned_road_id = _matched_alignment_road_id(
            road_alignments.get(row["case_id"])
        )
        result[row["case_id"]][row["agent_type"]].append(
            {
                "dsl_answer_id": row["dsl_answer_id"],
                "answer_set_id": row["answer_set_id"],
                "subtask_key": row["subtask_key"],
                "agent_type": row["agent_type"],
                "road_id": aligned_road_id or row["road_id"],
                "dsl": _synchronize_dsl_road_ids(
                    _load_json(row["dsl_json"], {}), aligned_road_id
                ),
                "review_status": row["review_status"],
                "metadata": _load_json(row["metadata_json"], {}),
                "tool_answers": [
                    _synchronize_tool_answer(answer, aligned_road_id)
                    for answer in tool_answers.get(row["dsl_answer_id"], [])
                ],
            }
        )
    return result


def _tool_answer_row(connection, tool_answer_id):
    row = connection.execute(
        """
        SELECT tool_call_answers.*, dsl_answers.agent_type
        FROM tool_call_answers
        JOIN dsl_answers USING (dsl_answer_id)
        WHERE tool_answer_id = ?
        """,
        (tool_answer_id,),
    ).fetchone()
    if row is None:
        raise ValueError(f"Tool answer does not exist: {tool_answer_id}")
    return row


def _tool_answer_dict(row):
    return {
        "tool_answer_id": row["tool_answer_id"],
        "dsl_answer_id": row["dsl_answer_id"],
        "variant_index": row["variant_index"],
        "tool_calls": _load_json(row["tool_call_json"], []),
        "resolved_result": _load_json(row["resolved_result_json"], None),
        "internal_commands": _load_json(row["internal_command_json"], None),
        "random_seed": row["random_seed"],
        "review_status": row["review_status"],
        "metadata": _load_json(row["metadata_json"], {}),
        "created_at": row["created_at"],
        "updated_at": row["updated_at"],
    }


def list_tool_call_answers(connection, dsl_answer_id, road_id=None):
    rows = connection.execute(
        """
        SELECT * FROM tool_call_answers
        WHERE dsl_answer_id = ?
        ORDER BY variant_index
        """,
        (dsl_answer_id,),
    ).fetchall()
    return [
        _synchronize_tool_answer(_tool_answer_dict(row), road_id)
        for row in rows
    ]


def _validate_tool_answer_content(
    agent_type,
    tool_calls,
    resolved_result,
    internal_commands,
    random_seed,
    metadata,
):
    if not isinstance(tool_calls, list):
        raise ValueError("Tool calls must be an array.")
    for index, call in enumerate(tool_calls, start=1):
        if (
            not isinstance(call, dict)
            or not isinstance(call.get("name"), str)
            or not isinstance(call.get("args"), dict)
        ):
            raise ValueError(
                f"Tool call {index} must contain string name and object args."
            )
    validate_resolved_result(agent_type, resolved_result)
    if agent_type == "curve" and tool_calls and resolved_result is None:
        raise ValueError(
            "Curve tool answers with calls require optimized control points."
        )
    if internal_commands is not None and (
        not isinstance(internal_commands, list)
        or any(not isinstance(command, str) for command in internal_commands)
    ):
        raise ValueError("Internal commands must be a string array.")
    if random_seed is not None and not isinstance(random_seed, int):
        raise ValueError("random_seed must be an integer or null.")
    if metadata is not None and not isinstance(metadata, dict):
        raise ValueError("metadata must be an object.")


def create_tool_call_answer(
    connection,
    dsl_answer_id,
    tool_calls,
    resolved_result=None,
    internal_commands=None,
    random_seed=None,
    metadata=None,
    review_status="approved",
):
    dsl_answer = connection.execute(
        "SELECT agent_type FROM dsl_answers WHERE dsl_answer_id = ?",
        (dsl_answer_id,),
    ).fetchone()
    if dsl_answer is None:
        raise ValueError(f"DSL answer does not exist: {dsl_answer_id}")
    _validate_tool_answer_content(
        dsl_answer["agent_type"],
        tool_calls,
        resolved_result,
        internal_commands,
        random_seed,
        metadata,
    )
    if review_status not in {"draft", "approved", "rejected"}:
        raise ValueError(f"Unsupported tool answer review status: {review_status}")
    variant_index = connection.execute(
        """
        SELECT COALESCE(MAX(variant_index), 0) + 1
        FROM tool_call_answers
        WHERE dsl_answer_id = ?
        """,
        (dsl_answer_id,),
    ).fetchone()[0]
    tool_answer_id = insert_tool_call_answer(
        connection,
        dsl_answer_id,
        variant_index,
        tool_calls,
        resolved_result=resolved_result,
        internal_commands=internal_commands,
        random_seed=random_seed,
        review_status=review_status,
        metadata=metadata,
    )
    return _tool_answer_dict(_tool_answer_row(connection, tool_answer_id))


def update_tool_call_answer(
    connection,
    tool_answer_id,
    tool_calls,
    resolved_result=None,
    internal_commands=None,
    random_seed=None,
    metadata=None,
    review_status="approved",
):
    current = _tool_answer_row(connection, tool_answer_id)
    _validate_tool_answer_content(
        current["agent_type"],
        tool_calls,
        resolved_result,
        internal_commands,
        random_seed,
        metadata,
    )
    if review_status not in {"draft", "approved", "rejected"}:
        raise ValueError(f"Unsupported tool answer review status: {review_status}")
    connection.execute(
        """
        UPDATE tool_call_answers
        SET tool_call_json = ?, resolved_result_json = ?,
            internal_command_json = ?, random_seed = ?, review_status = ?,
            metadata_json = ?,
            updated_at = strftime('%Y-%m-%dT%H:%M:%fZ', 'now')
        WHERE tool_answer_id = ?
        """,
        (
            _json_text(tool_calls),
            None if resolved_result is None else _json_text(resolved_result),
            None if internal_commands is None else _json_text(internal_commands),
            random_seed,
            review_status,
            _json_text(metadata or {}),
            tool_answer_id,
        ),
    )
    return _tool_answer_dict(_tool_answer_row(connection, tool_answer_id))


def update_tool_call_review_status(connection, tool_answer_id, review_status):
    if review_status not in {"draft", "approved", "rejected"}:
        raise ValueError(f"Unsupported tool answer review status: {review_status}")
    _tool_answer_row(connection, tool_answer_id)
    connection.execute(
        """
        UPDATE tool_call_answers
        SET review_status = ?,
            updated_at = strftime('%Y-%m-%dT%H:%M:%fZ', 'now')
        WHERE tool_answer_id = ?
        """,
        (review_status, tool_answer_id),
    )
    return _tool_answer_dict(_tool_answer_row(connection, tool_answer_id))


def delete_tool_call_answer(connection, tool_answer_id):
    _tool_answer_row(connection, tool_answer_id)
    connection.execute(
        "DELETE FROM tool_call_answers WHERE tool_answer_id = ?",
        (tool_answer_id,),
    )


def _case_review_dict(row):
    return {
        "review_id": row["review_id"],
        "case_id": row["case_id"],
        "run_id": row["run_id"],
        "task_id": row["task_id"],
        "decision": row["decision"],
        "note": row["note"],
        "review_source": row["review_source"],
        "trace_source": row["trace_source"],
        "metadata": _load_json(row["metadata_json"], {}),
        "created_at": row["created_at"],
        "updated_at": row["updated_at"],
    }


def list_case_reviews(connection, run_id=None):
    query = "SELECT * FROM case_reviews"
    parameters = ()
    if run_id is not None:
        query += " WHERE run_id = ?"
        parameters = (str(run_id),)
    query += " ORDER BY run_id, task_id, review_id"
    return [
        _case_review_dict(row)
        for row in connection.execute(query, parameters).fetchall()
    ]


def upsert_case_review(
    connection,
    case_id,
    run_id,
    task_id,
    decision,
    note="",
    review_source="human",
    trace_source=None,
    metadata=None,
):
    _find_case(connection, case_id=case_id)
    run_id = str(run_id or "").strip()
    task_id = str(task_id or "").strip()
    if not run_id:
        raise ValueError("Review run_id must not be empty.")
    if not task_id:
        raise ValueError("Review task_id must not be empty.")
    if decision not in {"pending", "pass", "question", "fail"}:
        raise ValueError(f"Unsupported review decision: {decision}")
    if review_source not in {"human", "auto"}:
        raise ValueError(f"Unsupported review source: {review_source}")
    if not isinstance(metadata, (dict, type(None))):
        raise ValueError("Review metadata must be an object.")

    current = connection.execute(
        """
        SELECT * FROM case_reviews
        WHERE case_id = ? AND run_id = ? AND task_id = ?
        """,
        (case_id, run_id, task_id),
    ).fetchone()
    if current is not None and current["review_source"] == "human" and review_source == "auto":
        return _case_review_dict(current)

    connection.execute(
        """
        INSERT INTO case_reviews (
            case_id, run_id, task_id, decision, note, review_source,
            trace_source, metadata_json, created_at, updated_at
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?,
                  strftime('%Y-%m-%dT%H:%M:%fZ', 'now'),
                  strftime('%Y-%m-%dT%H:%M:%fZ', 'now'))
        ON CONFLICT(case_id, run_id, task_id) DO UPDATE SET
            decision = excluded.decision,
            note = excluded.note,
            review_source = excluded.review_source,
            trace_source = excluded.trace_source,
            metadata_json = excluded.metadata_json,
            updated_at = strftime('%Y-%m-%dT%H:%M:%fZ', 'now')
        """,
        (
            case_id,
            run_id,
            task_id,
            decision,
            str(note or ""),
            review_source,
            None if trace_source is None else str(trace_source),
            _json_text(metadata or {}),
        ),
    )
    stored = connection.execute(
        """
        SELECT * FROM case_reviews
        WHERE case_id = ? AND run_id = ? AND task_id = ?
        """,
        (case_id, run_id, task_id),
    ).fetchone()
    return _case_review_dict(stored)


def _load_json(value, fallback):
    if value is None or value == "":
        return fallback
    return json.loads(value)


def list_request_rows(connection):
    rows = []
    for row in connection.execute(
        "SELECT * FROM test_cases ORDER BY source_row"
    ):
        road_context = _load_json(row["road_context_json"], {})
        rows.append(
            {
                "case_id": row["case_id"],
                "source_row": row["source_row"],
                "输入": row["input_text"],
                "算子": row["operator_text"],
                "类型": row["case_type"],
                "width": road_context.get("width"),
                "slope": road_context.get("slope"),
                "curve": road_context.get("curve"),
            }
        )
    return rows


def load_request_case_map(database_path=DEFAULT_DATABASE_PATH):
    connection = connect_database(database_path)
    try:
        return {row["输入"].strip(): row for row in list_request_rows(connection)}
    finally:
        connection.close()


def load_test_case_road_alignment_map(database_path=DEFAULT_DATABASE_PATH):
    """Return the read-only real-scene alignment indexed by source input text."""
    database_uri = Path(database_path).resolve().as_uri() + "?mode=ro"
    connection = sqlite3.connect(database_uri, uri=True)
    connection.row_factory = sqlite3.Row
    try:
        table_exists = connection.execute(
            """
            SELECT 1
            FROM sqlite_master
            WHERE type = 'table' AND name = 'test_case_road_alignments'
            """
        ).fetchone()
        if table_exists is None:
            raise RuntimeError(
                "Database is missing test_case_road_alignments."
            )

        rows = connection.execute(
            """
            SELECT
                test_cases.input_text,
                alignments.case_id,
                alignments.match_status,
                alignments.scene_file_name,
                alignments.map_file_name,
                alignments.road_id,
                alignments.s_start,
                alignments.s_end,
                alignments.route_lane_ids_json,
                alignments.driving_lane_ids_json
            FROM test_case_road_alignments AS alignments
            JOIN test_cases ON test_cases.case_id = alignments.case_id
            ORDER BY test_cases.source_row
            """
        ).fetchall()
        return {
            row["input_text"].strip(): {
                "case_id": row["case_id"],
                "match_status": row["match_status"],
                "scene_file_name": row["scene_file_name"],
                "map_file_name": row["map_file_name"],
                "road_id": row["road_id"],
                "s_start": row["s_start"],
                "s_end": row["s_end"],
                "route_lane_ids": _load_json(row["route_lane_ids_json"], []),
                "driving_lane_ids": _load_json(
                    row["driving_lane_ids_json"], []
                ),
            }
            for row in rows
        }
    finally:
        connection.close()


def _find_case(connection, *, case_id=None, input_text=None):
    if case_id is not None:
        row = connection.execute(
            "SELECT * FROM test_cases WHERE case_id = ?", (case_id,)
        ).fetchone()
    elif input_text is not None:
        row = connection.execute(
            "SELECT * FROM test_cases WHERE input_text = ?", (str(input_text),)
        ).fetchone()
    else:
        raise ValueError("case_id or input_text is required.")
    if row is None:
        identifier = case_id if case_id is not None else input_text
        raise ValueError(f"Request test case does not exist: {identifier}")
    return row


def update_case_input(connection, case_id, input_text):
    input_text = str(input_text).strip()
    if not input_text:
        raise ValueError("Input text must not be empty.")
    _find_case(connection, case_id=case_id)
    try:
        connection.execute(
            """
            UPDATE test_cases
            SET input_text = ?, updated_at = strftime('%Y-%m-%dT%H:%M:%fZ', 'now')
            WHERE case_id = ?
            """,
            (input_text, case_id),
        )
    except sqlite3.IntegrityError as exc:
        raise ValueError(f"Input text already exists: {input_text}") from exc


def save_road_context_values(connection, input_text, values_by_type, append=False):
    row = _find_case(connection, input_text=input_text)
    context = _load_json(row["road_context_json"], {})
    for agent_type, values in values_by_type.items():
        if agent_type not in {"slope", "curve"}:
            raise ValueError(f"Unsupported road context type: {agent_type}")
        normalized = [str(value).strip() for value in values if str(value).strip()]
        if not normalized:
            continue
        if append:
            current = [
                item.strip()
                for item in str(context.get(agent_type) or "").splitlines()
                if item.strip()
            ]
            normalized = current + normalized
        context[agent_type] = "\n".join(normalized)
    connection.execute(
        """
        UPDATE test_cases
        SET road_context_json = ?,
            updated_at = strftime('%Y-%m-%dT%H:%M:%fZ', 'now')
        WHERE case_id = ?
        """,
        (_json_text(context), row["case_id"]),
    )


def _merge_dsl_answer_rows(rows):
    grouped = {}
    for row in rows:
        key = row["answer_set_id"]
        if key not in grouped:
            grouped[key] = {
                "answer_set_id": key,
                "first_id": row["dsl_answer_id"],
                "payload": None,
                "operations": [],
                "metadata": _load_json(row["metadata_json"], {}),
                "review_status": row["review_status"],
            }
        item = grouped[key]
        payload = _load_json(row["dsl_json"], {})
        if item["payload"] is None:
            item["payload"] = {
                key: value for key, value in payload.items() if key != "operations"
            }
        item["operations"].extend(payload.get("operations", []))
        if row["review_status"] == "approved":
            item["review_status"] = "approved"

    entries = []
    for item in sorted(grouped.values(), key=lambda value: value["first_id"]):
        payload = dict(item["payload"] or {})
        payload["operations"] = item["operations"]
        entry = dict(item["metadata"])
        entry["value"] = payload
        if item["review_status"] is not None:
            entry["review_status"] = item["review_status"]
        entries.append(entry)
    return entries


def list_dsl_answer_entries(connection, case_id, agent_type):
    if agent_type not in AGENT_TYPES:
        raise ValueError(f"Unsupported agent type: {agent_type}")
    _find_case(connection, case_id=case_id)
    rows = connection.execute(
        """
        SELECT *
        FROM dsl_answers
        WHERE case_id = ? AND agent_type = ?
        ORDER BY dsl_answer_id
        """,
        (case_id, agent_type),
    ).fetchall()
    alignment = list_case_road_alignments(connection, [case_id]).get(case_id)
    return _synchronize_dsl_entries(
        _merge_dsl_answer_rows(rows), _matched_alignment_road_id(alignment)
    )


def _normalize_dsl_entries(entries, agent_type):
    if not isinstance(entries, list):
        raise ValueError("DSL answer entries must be an array.")
    normalized = []
    for index, entry in enumerate(entries, start=1):
        if not isinstance(entry, dict):
            raise ValueError(f"DSL answer entry {index} must be an object.")
        payload = entry.get("value", entry)
        if isinstance(payload, str):
            payload = json.loads(payload)
        if not isinstance(payload, dict) or not isinstance(payload.get("operations"), list):
            raise ValueError(f"DSL answer entry {index} must contain operations.")
        if any(
            not isinstance(operation, dict)
            or operation.get("operation") != agent_type
            for operation in payload["operations"]
        ):
            raise ValueError(
                f"DSL answer entry {index} contains a non-{agent_type} operation."
            )
        metadata = {
            key: value
            for key, value in entry.items()
            if key not in {"value", "review_status"}
        }
        review_status = entry.get("review_status")
        if review_status not in {None, "draft", "approved", "rejected", "stale"}:
            raise ValueError(f"Unsupported review status: {review_status}")
        normalized.append((payload, metadata, review_status))
    return normalized


def replace_dsl_answer_entries(connection, case_id, agent_type, entries):
    if agent_type not in AGENT_TYPES:
        raise ValueError(f"Unsupported agent type: {agent_type}")
    case = _find_case(connection, case_id=case_id)
    normalized = _normalize_dsl_entries(entries, agent_type)
    alignment = list_case_road_alignments(connection, [case_id]).get(case_id)
    aligned_road_id = _matched_alignment_road_id(alignment)
    if aligned_road_id:
        normalized = [
            (
                _synchronize_dsl_road_ids(payload, aligned_road_id),
                metadata,
                review_status,
            )
            for payload, metadata, review_status in normalized
        ]
    existing_rows = connection.execute(
        """
        SELECT * FROM dsl_answers
        WHERE case_id = ? AND agent_type = ?
        """,
        (case_id, agent_type),
    ).fetchall()
    existing = {
        (row["answer_set_id"], row["subtask_key"]): row
        for row in existing_rows
    }
    retained_ids = set()

    for candidate_index, (payload, metadata, review_status) in enumerate(
        normalized, start=1
    ):
        answer_set_id = f"{case_id}_set_{candidate_index}"
        for subtask_index, (road_id, grouped_payload) in enumerate(
            _group_dsl_by_road(payload), start=1
        ):
            subtask_key = f"{agent_type}_{subtask_index}"
            key = (answer_set_id, subtask_key)
            current = existing.get(key)
            dsl_json = _json_text(grouped_payload)
            metadata_json = _json_text(metadata)
            if current is None:
                cursor = connection.execute(
                    """
                    INSERT INTO dsl_answers (
                        answer_set_id, case_id, subtask_key, agent_type, road_id,
                        subtask_description, dsl_json, review_status, metadata_json,
                        created_at, updated_at
                    ) VALUES (?, ?, ?, ?, ?, NULL, ?, ?, ?,
                              strftime('%Y-%m-%dT%H:%M:%fZ', 'now'),
                              strftime('%Y-%m-%dT%H:%M:%fZ', 'now'))
                    """,
                    (
                        answer_set_id,
                        case_id,
                        subtask_key,
                        agent_type,
                        road_id,
                        dsl_json,
                        review_status,
                        metadata_json,
                    ),
                )
                retained_ids.add(cursor.lastrowid)
                continue

            retained_ids.add(current["dsl_answer_id"])
            dsl_changed = current["dsl_json"] != dsl_json or current["road_id"] != road_id
            next_status = (
                current["review_status"] if review_status is None else review_status
            )
            connection.execute(
                """
                UPDATE dsl_answers
                SET road_id = ?, dsl_json = ?, review_status = ?, metadata_json = ?,
                    updated_at = strftime('%Y-%m-%dT%H:%M:%fZ', 'now')
                WHERE dsl_answer_id = ?
                """,
                (
                    road_id,
                    dsl_json,
                    next_status,
                    metadata_json,
                    current["dsl_answer_id"],
                ),
            )
            if dsl_changed:
                connection.execute(
                    """
                    UPDATE tool_call_answers
                    SET review_status = 'stale',
                        updated_at = strftime('%Y-%m-%dT%H:%M:%fZ', 'now')
                    WHERE dsl_answer_id = ?
                    """,
                    (current["dsl_answer_id"],),
                )

    obsolete_ids = [
        row["dsl_answer_id"]
        for row in existing_rows
        if row["dsl_answer_id"] not in retained_ids
    ]
    if obsolete_ids:
        placeholders = ",".join("?" for _ in obsolete_ids)
        connection.execute(
            f"DELETE FROM dsl_answers WHERE dsl_answer_id IN ({placeholders})",
            obsolete_ids,
        )
    return list_dsl_answer_entries(connection, case["case_id"], agent_type)


def build_viewer_request_rows(connection):
    headers = [
        "case_id",
        "输入",
        "算子",
        "类型",
        "width",
        "slope",
        "curve",
        "width answer",
        "slope answer",
        "curve answer",
    ]
    rows = [headers]
    for case in list_request_rows(connection):
        answer_cells = {}
        for agent_type, column in ANSWER_COLUMNS.items():
            entries = list_dsl_answer_entries(
                connection, case["case_id"], agent_type
            )
            answer_cells[column] = "" if not entries else _json_text(entries)
        rows.append([case.get(header, answer_cells.get(header)) for header in headers])
    return rows


def verify_sheet1_parity(workbook_path, database_path):
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    connection = connect_database(database_path)
    try:
        if "Sheet1" not in workbook.sheetnames:
            raise ValueError("Workbook must contain Sheet1.")
        worksheet = workbook["Sheet1"]
        header_row = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True))
        headers = {value: index for index, value in enumerate(header_row) if value}
        missing = [column for column in REQUIRED_COLUMNS if column not in headers]
        if missing:
            raise ValueError(f"Sheet1 is missing columns: {', '.join(missing)}")

        expected_cases = []
        expected_answers = []
        answer_source_cells = 0
        case_count = 0
        for source_row, row in enumerate(
            worksheet.iter_rows(min_row=2, values_only=True), start=2
        ):
            row_data = {
                header: row[index] if index < len(row) else None
                for header, index in headers.items()
            }
            if row_data.get("输入") is None or not str(row_data["输入"]).strip():
                continue
            case_count += 1
            case_id = f"req_{case_count:06d}"
            expected_cases.append(
                {
                    "case_id": case_id,
                    "source_row": source_row,
                    "input_text": str(row_data["输入"]).strip(),
                    "operator_text": row_data.get("算子"),
                    "operator_types_json": _json_text(
                        _parse_operator_types(row_data.get("算子"))
                    ),
                    "case_type": row_data.get("类型"),
                    "road_context_json": _json_text(
                        {
                            "width": row_data.get("width"),
                            "slope": row_data.get("slope"),
                            "curve": row_data.get("curve"),
                        }
                    ),
                    "created_at": None,
                    "updated_at": None,
                }
            )
            records, source_cells = _build_dsl_answer_records(
                case_id, source_row, row_data
            )
            expected_answers.extend(records)
            answer_source_cells += source_cells

        actual_cases = [
            dict(row)
            for row in connection.execute(
                "SELECT * FROM test_cases ORDER BY source_row"
            )
        ]
        actual_answers = [
            {
                key: row[key]
                for key in (
                    "answer_set_id",
                    "case_id",
                    "subtask_key",
                    "agent_type",
                    "road_id",
                    "subtask_description",
                    "dsl_json",
                    "review_status",
                    "metadata_json",
                    "created_at",
                    "updated_at",
                )
            }
            for row in connection.execute(
                "SELECT * FROM dsl_answers ORDER BY dsl_answer_id"
            )
        ]
        if actual_cases != expected_cases:
            raise RuntimeError("test_cases does not exactly match Sheet1.")
        if actual_answers != expected_answers:
            raise RuntimeError("dsl_answers does not exactly match Sheet1 answers.")
        tool_answer_count = connection.execute(
            "SELECT COUNT(*) FROM tool_call_answers"
        ).fetchone()[0]
        if tool_answer_count != 0:
            raise RuntimeError("Sheet1 migration must not populate tool_call_answers.")
        return {
            "sheet": "Sheet1",
            "test_cases": len(expected_cases),
            "answer_source_cells": answer_source_cells,
            "dsl_answers": len(expected_answers),
            "tool_call_answers": tool_answer_count,
        }
    finally:
        connection.close()
        workbook.close()


def verify_database(connection):
    integrity = connection.execute("PRAGMA integrity_check").fetchone()[0]
    if integrity != "ok":
        raise RuntimeError(f"SQLite integrity check failed: {integrity}")
    foreign_key_errors = connection.execute("PRAGMA foreign_key_check").fetchall()
    if foreign_key_errors:
        raise RuntimeError("SQLite foreign key check failed.")
    tables = [
        row[0]
        for row in connection.execute(
            """
            SELECT name
            FROM sqlite_master
            WHERE type = 'table' AND name NOT LIKE 'sqlite_%'
            ORDER BY name
            """
        )
    ]
    required_tables = {
        "case_reviews",
        "dsl_answers",
        "test_cases",
        "tool_call_answers",
    }
    allowed_tables = required_tables | {"test_case_road_alignments"}
    if not required_tables.issubset(tables) or not set(tables).issubset(allowed_tables):
        raise RuntimeError(f"Unexpected application tables: {tables}")
    return {
        "schema_version": connection.execute("PRAGMA user_version").fetchone()[0],
        "test_cases": connection.execute(
            "SELECT COUNT(*) FROM test_cases"
        ).fetchone()[0],
        "dsl_answers": connection.execute(
            "SELECT COUNT(*) FROM dsl_answers"
        ).fetchone()[0],
        "tool_call_answers": connection.execute(
            "SELECT COUNT(*) FROM tool_call_answers"
        ).fetchone()[0],
        "case_reviews": connection.execute(
            "SELECT COUNT(*) FROM case_reviews"
        ).fetchone()[0],
    }


def _build_parser():
    parser = argparse.ArgumentParser(description="Manage the request test database.")
    subparsers = parser.add_subparsers(dest="command", required=True)

    migrate_parser = subparsers.add_parser(
        "migrate", description="Migrate only Sheet1 from 请求测试.xlsx."
    )
    migrate_parser.add_argument("--workbook", default=str(DEFAULT_WORKBOOK_PATH))
    migrate_parser.add_argument("--database", default=str(DEFAULT_DATABASE_PATH))
    migrate_parser.add_argument("--overwrite", action="store_true")

    verify_parser = subparsers.add_parser("verify", description="Verify the database.")
    verify_parser.add_argument("--database", default=str(DEFAULT_DATABASE_PATH))
    verify_parser.add_argument("--workbook", default=str(DEFAULT_WORKBOOK_PATH))
    return parser


def main():
    args = _build_parser().parse_args()
    if args.command == "migrate":
        result = migrate_sheet1(args.workbook, args.database, overwrite=args.overwrite)
    else:
        connection = connect_database(args.database)
        try:
            result = verify_database(connection)
        finally:
            connection.close()
        result["sheet1_parity"] = verify_sheet1_parity(
            args.workbook, args.database
        )
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
