"""SQLite persistence and spreadsheet export for EGO compliance analysis."""

from __future__ import annotations

from datetime import datetime, timezone
import json
from pathlib import Path
import sqlite3
from typing import Any, Mapping

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet


REVIEW_STATUSES = {"unreviewed", "confirmed", "dismissed"}
DEFAULT_EXPORT_DIRECTORY = Path(r"D:\download")
VIOLATION_EXPORT_TYPES = (
    ("left_line_crossing", "左侧越线"),
    ("right_line_crossing", "右侧越线"),
    ("longitudinal_slide", "纵向滑坡"),
    ("fall_off_road", "跌出路面"),
    ("collision_candidate", "碰撞候选"),
    ("vehicle_stuck", "车辆停滞"),
)


class EgoAnalysisStore:
    """Persist EGO batches separately from NPC visualization history."""

    def __init__(
        self,
        database_path: str | Path,
        export_directory: str | Path = DEFAULT_EXPORT_DIRECTORY,
    ):
        self.database_path = Path(database_path)
        self.export_directory = Path(export_directory)
        self.database_path.parent.mkdir(parents=True, exist_ok=True)
        with self._connect() as connection:
            connection.executescript(
                """
                CREATE TABLE IF NOT EXISTS ego_analysis_history (
                    id TEXT PRIMARY KEY,
                    source_name TEXT NOT NULL,
                    summary_json TEXT NOT NULL,
                    thresholds_json TEXT NOT NULL,
                    created_at TEXT NOT NULL,
                    updated_at TEXT NOT NULL
                );
                CREATE TABLE IF NOT EXISTS ego_analysis_runs (
                    batch_id TEXT NOT NULL,
                    run_index INTEGER NOT NULL,
                    json_name TEXT NOT NULL,
                    map_name TEXT NOT NULL,
                    sim_result TEXT NOT NULL,
                    payload_json TEXT NOT NULL,
                    collision_review_status TEXT NOT NULL DEFAULT 'unreviewed',
                    collision_review_note TEXT NOT NULL DEFAULT '',
                    PRIMARY KEY (batch_id, run_index),
                    FOREIGN KEY (batch_id) REFERENCES ego_analysis_history(id)
                        ON DELETE CASCADE
                );
                """
            )

    def save(
        self,
        batch_id: str,
        source_name: str,
        analysis: Mapping[str, object],
    ) -> dict[str, object]:
        now = datetime.now(timezone.utc).isoformat()
        summary_json = json.dumps(
            analysis["summary"],
            ensure_ascii=False,
            separators=(",", ":"),
        )
        thresholds_json = json.dumps(
            analysis["thresholds"],
            ensure_ascii=False,
            separators=(",", ":"),
        )
        with self._connect() as connection:
            connection.execute(
                """
                INSERT INTO ego_analysis_history (
                    id, source_name, summary_json, thresholds_json,
                    created_at, updated_at
                ) VALUES (?, ?, ?, ?, ?, ?)
                ON CONFLICT(id) DO UPDATE SET
                    source_name = excluded.source_name,
                    summary_json = excluded.summary_json,
                    thresholds_json = excluded.thresholds_json,
                    updated_at = excluded.updated_at
                """,
                (
                    batch_id,
                    source_name,
                    summary_json,
                    thresholds_json,
                    now,
                    now,
                ),
            )
            existing_reviews = {
                int(row[0]): (row[1], row[2])
                for row in connection.execute(
                    """
                    SELECT run_index, collision_review_status, collision_review_note
                    FROM ego_analysis_runs WHERE batch_id = ?
                    """,
                    (batch_id,),
                )
            }
            connection.execute(
                "DELETE FROM ego_analysis_runs WHERE batch_id = ?",
                (batch_id,),
            )
            for run in analysis["runs"]:
                run_index = int(run["runIndex"])
                default_review = run.get("collisionReview") or {}
                review = existing_reviews.get(
                    run_index,
                    (
                        str(default_review.get("status") or "unreviewed"),
                        str(default_review.get("note") or ""),
                    ),
                )
                connection.execute(
                    """
                    INSERT INTO ego_analysis_runs (
                        batch_id, run_index, json_name, map_name, sim_result,
                        payload_json, collision_review_status,
                        collision_review_note
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    (
                        batch_id,
                        run_index,
                        str(run["jsonName"]),
                        str(run["mapName"]),
                        str(run["simResult"]),
                        json.dumps(
                            run,
                            ensure_ascii=False,
                            separators=(",", ":"),
                        ),
                        review[0],
                        review[1],
                    ),
                )
        return self.get_item(batch_id)

    def list_items(self) -> list[dict[str, object]]:
        with self._connect() as connection:
            rows = connection.execute(
                """
                SELECT id, source_name, summary_json, created_at, updated_at
                FROM ego_analysis_history
                ORDER BY updated_at DESC
                """
            ).fetchall()
            return [
                self._history_item(connection, row)
                for row in rows
            ]

    def get_item(self, batch_id: str) -> dict[str, object]:
        with self._connect() as connection:
            row = connection.execute(
                """
                SELECT id, source_name, summary_json, created_at, updated_at
                FROM ego_analysis_history WHERE id = ?
                """,
                (batch_id,),
            ).fetchone()
            if row is None:
                raise KeyError(batch_id)
            return self._history_item(connection, row)

    def get_batch(self, batch_id: str) -> dict[str, object] | None:
        with self._connect() as connection:
            row = connection.execute(
                """
                SELECT source_name, summary_json, thresholds_json,
                       created_at, updated_at
                FROM ego_analysis_history WHERE id = ?
                """,
                (batch_id,),
            ).fetchone()
            if row is None:
                return None
            run_rows = connection.execute(
                """
                SELECT run_index, payload_json,
                       collision_review_status, collision_review_note
                FROM ego_analysis_runs
                WHERE batch_id = ?
                ORDER BY run_index
                """,
                (batch_id,),
            ).fetchall()
        runs = []
        for run_index, payload_json, review_status, review_note in run_rows:
            payload = json.loads(payload_json)
            runs.append(
                {
                    "runIndex": run_index,
                    "jsonName": payload["jsonName"],
                    "mapName": payload["mapName"],
                    "mapFile": payload["mapFile"],
                    "simResult": payload["simResult"],
                    "collisionCandidate": payload["collisionCandidate"],
                    "stuck": bool(payload.get("stuck", False)),
                    "collisionReview": {
                        "status": review_status,
                        "note": review_note,
                    },
                    "roadIds": payload["roadIds"],
                    "stats": payload["stats"],
                    "eventTypes": sorted(
                        {event["type"] for event in payload["events"]}
                    ),
                }
            )
        summary = self._reviewed_summary(json.loads(row[1]), runs)
        return {
            "id": batch_id,
            "sourceName": row[0],
            "summary": summary,
            "thresholds": json.loads(row[2]),
            "runs": runs,
            "createdAt": row[3],
            "updatedAt": row[4],
        }

    def get_run(self, batch_id: str, run_index: int) -> dict[str, object] | None:
        with self._connect() as connection:
            row = connection.execute(
                """
                SELECT payload_json, collision_review_status,
                       collision_review_note
                FROM ego_analysis_runs
                WHERE batch_id = ? AND run_index = ?
                """,
                (batch_id, run_index),
            ).fetchone()
        if row is None:
            return None
        payload = json.loads(row[0])
        payload["collisionReview"] = {
            "status": row[1],
            "note": row[2],
        }
        return payload

    def update_run_payload(
        self,
        batch_id: str,
        run_index: int,
        payload: Mapping[str, object],
    ) -> None:
        with self._connect() as connection:
            cursor = connection.execute(
                """
                UPDATE ego_analysis_runs
                SET payload_json = ?
                WHERE batch_id = ? AND run_index = ?
                """,
                (
                    json.dumps(
                        payload,
                        ensure_ascii=False,
                        separators=(",", ":"),
                    ),
                    batch_id,
                    run_index,
                ),
            )
        if cursor.rowcount == 0:
            raise KeyError((batch_id, run_index))

    def update_review(
        self,
        batch_id: str,
        run_index: int,
        status: str,
        note: str,
    ) -> dict[str, str]:
        if status not in REVIEW_STATUSES:
            raise ValueError(
                "Collision review status must be unreviewed, confirmed, or dismissed."
            )
        if len(note) > 2000:
            raise ValueError("Collision review note must be at most 2000 characters.")
        now = datetime.now(timezone.utc).isoformat()
        with self._connect() as connection:
            cursor = connection.execute(
                """
                UPDATE ego_analysis_runs
                SET collision_review_status = ?, collision_review_note = ?
                WHERE batch_id = ? AND run_index = ?
                """,
                (status, note, batch_id, run_index),
            )
            if cursor.rowcount == 0:
                raise KeyError((batch_id, run_index))
            connection.execute(
                "UPDATE ego_analysis_history SET updated_at = ? WHERE id = ?",
                (now, batch_id),
            )
        return {"status": status, "note": note}

    def delete(self, batch_id: str) -> bool:
        with self._connect() as connection:
            connection.execute(
                "DELETE FROM ego_analysis_runs WHERE batch_id = ?",
                (batch_id,),
            )
            cursor = connection.execute(
                "DELETE FROM ego_analysis_history WHERE id = ?",
                (batch_id,),
            )
        return cursor.rowcount > 0

    def export_xlsx(self, batch_id: str) -> str | None:
        batch = self.get_batch(batch_id)
        if batch is None:
            return None
        runs = [
            self.get_run(batch_id, int(run_meta["runIndex"]))
            for run_meta in batch["runs"]
        ]

        fieldnames = [
            "批次ID",
            "运行序号",
            "场景JSON",
            "地图",
            "仿真结果",
            "左侧越线次数",
            "左侧最大车身侵入距离(m)",
            "左侧越线道路/车道",
            "右侧越线次数",
            "右侧最大车身侵入距离(m)",
            "右侧越线道路/车道",
            "滑坡次数",
            "最大下坡路面位移(m)",
            "滑坡道路/车道",
            "跌落次数",
            "最高跌落速度(m/s)",
            "碰撞状态",
            "碰撞备注",
            "车辆停滞",
        ]
        workbook = Workbook()
        details_sheet = workbook.active
        details_sheet.title = "运行明细"
        details_sheet.append(fieldnames)
        for run in runs:
            events = list(run["events"])
            left_crossings = [
                event for event in events
                if event["type"] == "left_line_crossing"
            ]
            right_crossings = [
                event for event in events
                if event["type"] == "right_line_crossing"
            ]
            slides = [
                event for event in events
                if event["type"] == "longitudinal_slide"
            ]
            falls = [
                event for event in events
                if event["type"] == "fall_off_road"
            ]
            details_sheet.append(
                [
                    batch_id,
                    int(run["runIndex"]) + 1,
                    run["jsonName"],
                    run["mapName"],
                    run["simResult"],
                    len(left_crossings) or None,
                    _maximum_intrusion(left_crossings),
                    _event_road_lanes(left_crossings) or None,
                    len(right_crossings) or None,
                    _maximum_intrusion(right_crossings),
                    _event_road_lanes(right_crossings) or None,
                    len(slides) or None,
                    _maximum_event_value(slides, "downhillSurfaceDistance"),
                    _event_road_lanes(slides) or None,
                    len(falls) or None,
                    _maximum_event_value(falls, "peakSpeed"),
                    (
                        _collision_review_label(run["collisionReview"]["status"])
                        if run["collisionCandidate"]
                        else None
                    ),
                    (
                        run["collisionReview"]["note"]
                        if run["collisionCandidate"]
                        else None
                    ),
                    "是" if run.get("stuck", False) else None,
                ]
            )

        summary_sheet = workbook.create_sheet("违规汇总")
        summary_sheet.append(
            ["违规类型", "触发次数", "触发场景数", "触发场景占比"]
        )
        run_counts = [_violation_counts(run) for run in runs]
        for violation_type, label in VIOLATION_EXPORT_TYPES:
            trigger_count = sum(counts[violation_type] for counts in run_counts)
            scene_count = sum(
                counts[violation_type] > 0 for counts in run_counts
            )
            summary_sheet.append(
                [
                    label,
                    trigger_count,
                    scene_count,
                    scene_count / len(runs) if runs else 0,
                ]
            )
            summary_sheet.cell(summary_sheet.max_row, 4).number_format = "0.0%"
        matrix_sheet = workbook.create_sheet("场景违规矩阵")
        matrix_sheet.append(
            [
                "运行序号",
                "场景JSON",
                "地图",
                "仿真结果",
                *[label for _, label in VIOLATION_EXPORT_TYPES],
            ]
        )
        for run, counts in zip(runs, run_counts):
            matrix_sheet.append(
                [
                    int(run["runIndex"]) + 1,
                    run["jsonName"],
                    run["mapName"],
                    run["simResult"],
                    *[
                        "是" if counts[violation_type] > 0 else "否"
                        for violation_type, _ in VIOLATION_EXPORT_TYPES
                    ],
                ]
            )
        _format_export_sheet(
            details_sheet,
            [
                68, 10, 34, 18, 28, 14, 24, 24, 14, 24,
                24, 12, 22, 22, 12, 20, 14, 32, 12,
            ],
        )
        _format_export_sheet(summary_sheet, [18, 14, 16, 18])
        _format_export_sheet(
            matrix_sheet,
            [10, 34, 18, 28, 14, 14, 14, 14, 14, 14],
        )
        self.export_directory.mkdir(parents=True, exist_ok=True)
        file_name = f"{Path(batch['sourceName']).stem}_ego_analysis.xlsx"
        export_path = self.export_directory / file_name
        if export_path.exists():
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
            file_name = (
                f"{Path(batch['sourceName']).stem}_ego_analysis_{timestamp}.xlsx"
            )
            export_path = self.export_directory / file_name
        workbook.save(export_path)
        return file_name

    def _connect(self) -> sqlite3.Connection:
        connection = sqlite3.connect(self.database_path, timeout=30)
        connection.execute("PRAGMA foreign_keys = ON")
        return connection

    @staticmethod
    def _history_item(
        connection: sqlite3.Connection,
        row: tuple[Any, ...],
    ) -> dict[str, object]:
        run_rows = connection.execute(
            """
            SELECT payload_json, collision_review_status, collision_review_note
            FROM ego_analysis_runs WHERE batch_id = ?
            """,
            (row[0],),
        ).fetchall()
        runs = []
        for payload_json, status, note in run_rows:
            payload = json.loads(payload_json)
            runs.append(
                {
                    "collisionCandidate": payload["collisionCandidate"],
                    "collisionReview": {"status": status, "note": note},
                }
            )
        summary = EgoAnalysisStore._reviewed_summary(json.loads(row[2]), runs)
        return {
            "id": row[0],
            "sourceName": row[1],
            "summary": summary,
            "createdAt": row[3],
            "updatedAt": row[4],
        }

    @staticmethod
    def _reviewed_summary(
        summary: dict[str, object],
        runs: list[Mapping[str, object]],
    ) -> dict[str, object]:
        result = dict(summary)
        candidate_reviews = [
            run["collisionReview"]["status"]
            for run in runs
            if run["collisionCandidate"]
        ]
        result["unreviewedCollisionCount"] = candidate_reviews.count("unreviewed")
        result["confirmedCollisionCount"] = candidate_reviews.count("confirmed")
        result["dismissedCollisionCount"] = candidate_reviews.count("dismissed")
        return result


def _maximum_intrusion(events: list[Mapping[str, object]]) -> float | None:
    clearances = [
        float(event["minimumClearance"])
        for event in events
        if event.get("minimumClearance") not in (None, "")
    ]
    if not clearances:
        return None
    return max(max(0.0, 0.95 - clearance) for clearance in clearances)


def _maximum_event_value(
    events: list[Mapping[str, object]],
    key: str,
) -> float | None:
    values = [
        float(event[key])
        for event in events
        if event.get(key) not in (None, "")
    ]
    return max(values) if values else None


def _event_road_lanes(events: list[Mapping[str, object]]) -> str:
    road_lanes: list[str] = []
    for event in events:
        road_id = event.get("roadId")
        lane_id = event.get("laneId")
        if road_id in (None, "") and lane_id in (None, ""):
            continue
        if road_id not in (None, "") and lane_id not in (None, ""):
            value = f"R{road_id}/L{lane_id}"
        elif road_id not in (None, ""):
            value = f"R{road_id}"
        else:
            value = f"L{lane_id}"
        if value not in road_lanes:
            road_lanes.append(value)
    return "；".join(road_lanes)


def _violation_counts(run: Mapping[str, object]) -> dict[str, int]:
    counts = {violation_type: 0 for violation_type, _ in VIOLATION_EXPORT_TYPES}
    for event in run["events"]:
        violation_type = event["type"]
        if violation_type in counts:
            counts[violation_type] += 1
    counts["collision_candidate"] = int(bool(run["collisionCandidate"]))
    counts["vehicle_stuck"] = int(bool(run.get("stuck", False)))
    return counts


def _format_export_sheet(sheet: Worksheet, widths: list[int]) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    header_fill = PatternFill("solid", fgColor="2F6B55")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 24
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def _collision_review_label(status: object) -> str:
    return {
        "unreviewed": "待复核",
        "confirmed": "已确认",
        "dismissed": "已排除",
    }.get(str(status), str(status))
