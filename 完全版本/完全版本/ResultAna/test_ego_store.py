from copy import deepcopy
from pathlib import Path
import tempfile
import unittest

from openpyxl import load_workbook

from visualization.ego_store import EgoAnalysisStore


def analysis_payload():
    point = {
        "ordinalNumber": 12,
        "speed": 4.5,
        "roadId": "8",
        "laneId": 1,
    }
    run = {
        "runIndex": 0,
        "jsonName": "mutated.json",
        "mapName": "slope",
        "mapFile": "slope.xodr",
        "simResult": "FAILURE: collision",
        "collisionCandidate": True,
        "collisionReview": {"status": "unreviewed", "note": ""},
        "roadIds": ["8"],
        "points": [point],
        "events": [
            {
                "type": "left_line_crossing",
                "side": "left",
                "startOrdinal": 12,
                "endOrdinal": 12,
                "pointCount": 1,
                "roadId": "8",
                "laneId": 1,
                "minimumClearance": 0.8,
                "peakSpeed": 4.5,
                "downhillSurfaceDistance": None,
                "evidence": "trajectory_geometry",
            }
        ],
        "stats": {
            "waypointCount": 1,
            "lineCrossingPointCount": 1,
            "slideEventCount": 0,
            "fallPointCount": 0,
        },
        "roads": [],
        "npcs": [],
        "summary": {"roadCount": 0, "npcCount": 1, "waypointCount": 1},
    }
    return {
        "summary": {
            "runCount": 1,
            "waypointCount": 1,
            "violationRunCount": 1,
            "lineCrossingRunCount": 1,
            "slideRunCount": 0,
            "fallRunCount": 0,
            "collisionCandidateCount": 1,
            "unreviewedCollisionCount": 1,
            "stuckRunCount": 0,
            "analysisErrorCount": 0,
        },
        "thresholds": {"vehicleWidth": 1.9},
        "runs": [run],
    }


class EgoAnalysisStoreTests(unittest.TestCase):
    def setUp(self):
        self.temporary_directory = tempfile.TemporaryDirectory()
        self.database_path = Path(self.temporary_directory.name) / "history.sqlite3"
        self.export_directory = Path(self.temporary_directory.name) / "exports"
        self.store = EgoAnalysisStore(
            self.database_path,
            export_directory=self.export_directory,
        )

    def tearDown(self):
        self.temporary_directory.cleanup()

    def test_persists_batches_runs_and_collision_review(self):
        self.store.save("batch", "ego-log.json", analysis_payload())

        self.assertEqual(self.store.list_items()[0]["summary"]["runCount"], 1)
        run = self.store.get_run("batch", 0)
        self.assertEqual(run["jsonName"], "mutated.json")
        self.assertEqual(run["collisionReview"]["status"], "unreviewed")

        self.store.update_review("batch", 0, "confirmed", "人工检查录像")
        reopened = EgoAnalysisStore(self.database_path)
        batch = reopened.get_batch("batch")
        run = reopened.get_run("batch", 0)

        self.assertEqual(batch["summary"]["confirmedCollisionCount"], 1)
        self.assertEqual(batch["summary"]["unreviewedCollisionCount"], 0)
        self.assertEqual(run["collisionReview"]["note"], "人工检查录像")

    def test_xlsx_exports_three_sheets_and_preserves_run_details(self):
        payload = analysis_payload()
        payload["runs"][0]["stuck"] = True
        payload["runs"][0]["events"].extend(
            [
                {
                    "type": "left_line_crossing",
                    "roadId": "9",
                    "laneId": 2,
                    "minimumClearance": 0.6,
                    "peakSpeed": 6.0,
                },
                {
                    "type": "right_line_crossing",
                    "roadId": "8",
                    "laneId": 1,
                    "minimumClearance": 0.7,
                    "peakSpeed": 5.0,
                },
                {
                    "type": "longitudinal_slide",
                    "roadId": "8",
                    "laneId": 1,
                    "downhillSurfaceDistance": 0.72,
                    "peakSpeed": 3.0,
                },
                {
                    "type": "longitudinal_slide",
                    "roadId": "9",
                    "laneId": 2,
                    "downhillSurfaceDistance": 0.91,
                    "peakSpeed": 3.5,
                },
                {
                    "type": "fall_off_road",
                    "roadId": None,
                    "laneId": None,
                    "peakSpeed": 51.2,
                },
            ]
        )
        self.store.save("batch", "ego-log.json", payload)
        self.store.update_review("batch", 0, "confirmed", "已确认")

        file_name = self.store.export_xlsx("batch")
        workbook = load_workbook(self.export_directory / file_name)

        self.assertEqual(file_name, "ego-log_ego_analysis.xlsx")
        self.assertEqual(workbook.sheetnames, ["运行明细", "违规汇总", "场景违规矩阵"])
        details = list(workbook["运行明细"].values)
        row = dict(zip(details[0], details[1]))
        self.assertEqual(row["左侧越线次数"], 2)
        self.assertEqual(row["左侧最大车身侵入距离(m)"], 0.35)
        self.assertEqual(row["左侧越线道路/车道"], "R8/L1；R9/L2")
        self.assertEqual(row["右侧越线次数"], 1)
        self.assertEqual(row["右侧最大车身侵入距离(m)"], 0.25)
        self.assertEqual(row["右侧越线道路/车道"], "R8/L1")
        self.assertEqual(row["滑坡次数"], 2)
        self.assertEqual(row["最大下坡路面位移(m)"], 0.91)
        self.assertEqual(row["滑坡道路/车道"], "R8/L1；R9/L2")
        self.assertEqual(row["跌落次数"], 1)
        self.assertEqual(row["最高跌落速度(m/s)"], 51.2)
        self.assertEqual(row["碰撞状态"], "已确认")
        self.assertEqual(row["碰撞备注"], "已确认")
        self.assertEqual(row["车辆停滞"], "是")

    def test_xlsx_summarizes_event_counts_scene_counts_and_ratios(self):
        payload = analysis_payload()
        violating_run = payload["runs"][0]
        violating_run["stuck"] = True
        violating_run["events"].extend(
            [
                {"type": "left_line_crossing"},
                {"type": "right_line_crossing"},
                {"type": "longitudinal_slide"},
                {"type": "longitudinal_slide"},
                {"type": "fall_off_road"},
            ]
        )
        clean_run = deepcopy(violating_run)
        clean_run.update(
            {
                "runIndex": 1,
                "jsonName": "clean.json",
                "simResult": "SUCCESS",
                "collisionCandidate": False,
                "stuck": False,
                "events": [],
            }
        )
        payload["runs"].append(clean_run)
        self.store.save("batch", "ego-log.json", payload)

        file_name = self.store.export_xlsx("batch")
        workbook = load_workbook(self.export_directory / file_name)
        rows = list(workbook["违规汇总"].values)

        self.assertEqual(
            rows[0],
            ("违规类型", "触发次数", "触发场景数", "触发场景占比"),
        )
        self.assertEqual(
            rows[1:],
            [
                ("左侧越线", 2, 1, 0.5),
                ("右侧越线", 1, 1, 0.5),
                ("纵向滑坡", 2, 1, 0.5),
                ("跌出路面", 1, 1, 0.5),
                ("碰撞候选", 1, 1, 0.5),
                ("车辆停滞", 1, 1, 0.5),
            ],
        )
        for cell in workbook["违规汇总"]["D"][1:]:
            self.assertEqual(cell.number_format, "0.0%")

    def test_xlsx_marks_each_scene_violation_with_yes_or_no(self):
        payload = analysis_payload()
        violating_run = payload["runs"][0]
        violating_run["stuck"] = True
        clean_run = deepcopy(violating_run)
        clean_run.update(
            {
                "runIndex": 1,
                "jsonName": "clean.json",
                "simResult": "SUCCESS",
                "collisionCandidate": False,
                "stuck": False,
                "events": [],
            }
        )
        payload["runs"].append(clean_run)
        self.store.save("batch", "ego-log.json", payload)

        file_name = self.store.export_xlsx("batch")
        workbook = load_workbook(self.export_directory / file_name)
        rows = list(workbook["场景违规矩阵"].values)

        self.assertEqual(
            rows[0],
            (
                "运行序号",
                "场景JSON",
                "地图",
                "仿真结果",
                "左侧越线",
                "右侧越线",
                "纵向滑坡",
                "跌出路面",
                "碰撞候选",
                "车辆停滞",
            ),
        )
        self.assertEqual(
            rows[1],
            (1, "mutated.json", "slope", "FAILURE: collision", "是", "否", "否", "否", "是", "是"),
        )
        self.assertEqual(
            rows[2],
            (2, "clean.json", "slope", "SUCCESS", "否", "否", "否", "否", "否", "否"),
        )

    def test_xlsx_keeps_existing_export_and_uses_a_timestamped_name(self):
        self.store.save("batch", "ego-log.json", analysis_payload())
        self.export_directory.mkdir()
        existing_path = self.export_directory / "ego-log_ego_analysis.xlsx"
        existing_path.write_bytes(b"existing")

        file_name = self.store.export_xlsx("batch")

        self.assertRegex(
            file_name,
            r"^ego-log_ego_analysis_\d{8}_\d{6}_\d{6}\.xlsx$",
        )
        self.assertEqual(existing_path.read_bytes(), b"existing")
        self.assertEqual(
            load_workbook(self.export_directory / file_name).sheetnames,
            ["运行明细", "违规汇总", "场景违规矩阵"],
        )

    def test_xlsx_leaves_absent_error_columns_blank(self):
        payload = analysis_payload()
        run = payload["runs"][0]
        run["simResult"] = "SUCCESS"
        run["collisionCandidate"] = False
        run["collisionReview"] = {"status": "unreviewed", "note": "旧备注"}
        run["events"] = []
        run["stuck"] = False
        self.store.save("batch", "ego-log.json", payload)

        file_name = self.store.export_xlsx("batch")
        rows = list(load_workbook(self.export_directory / file_name)["运行明细"].values)
        row = dict(zip(rows[0], rows[1]))

        error_columns = [
            "左侧越线次数",
            "左侧最大车身侵入距离(m)",
            "左侧越线道路/车道",
            "右侧越线次数",
            "右侧最大车身侵入距离(m)",
            "右侧越线道路/车道",
            "滑坡次数",
            "最大下坡路面位移(m)",
            "滑坡道路/车道",
            "跌落次数",
            "最高跌落速度(m/s)",
            "碰撞状态",
            "碰撞备注",
            "车辆停滞",
        ]
        self.assertEqual(len(rows), 2)
        self.assertTrue(all(row[column] is None for column in error_columns))

    def test_delete_removes_batch_and_runs(self):
        self.store.save("batch", "ego-log.json", analysis_payload())

        self.assertTrue(self.store.delete("batch"))
        self.assertIsNone(self.store.get_batch("batch"))
        self.assertIsNone(self.store.get_run("batch", 0))

    def test_rejects_invalid_review_status(self):
        self.store.save("batch", "ego-log.json", analysis_payload())

        with self.assertRaisesRegex(ValueError, "Collision review status"):
            self.store.update_review("batch", 0, "maybe", "")


if __name__ == "__main__":
    unittest.main()
