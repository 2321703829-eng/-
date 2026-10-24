from http.client import HTTPConnection
import json
from pathlib import Path
import tempfile
from threading import Thread
import unittest

from openpyxl import load_workbook

from test_ego_store import analysis_payload
from visualization.server import create_server


class EgoExportServerTests(unittest.TestCase):
    def setUp(self):
        self.temporary_directory = tempfile.TemporaryDirectory()
        self.directory = Path(self.temporary_directory.name)
        self.export_directory = self.directory / "exports"
        self.server = create_server(
            database_path=self.directory / "history.sqlite3",
            port=0,
            map_directory=self.directory / "maps",
            scenario_directory=self.directory / "scenarios",
            export_directory=self.export_directory,
        )
        self.server.RequestHandlerClass.ego_store.save(
            "batch",
            "ego-log.json",
            analysis_payload(),
        )
        self.thread = Thread(target=self.server.serve_forever, daemon=True)
        self.thread.start()

    def tearDown(self):
        self.server.shutdown()
        self.server.server_close()
        self.thread.join(timeout=5)
        self.temporary_directory.cleanup()

    def test_export_saves_xlsx_and_returns_json_without_attachment(self):
        connection = HTTPConnection(
            "127.0.0.1",
            self.server.server_address[1],
            timeout=10,
        )
        connection.request("POST", "/api/ego/batch/export")
        response = connection.getresponse()
        response_body = json.loads(response.read().decode("utf-8"))
        content_disposition = response.getheader("Content-Disposition")
        content_type = response.getheader("Content-Type")
        connection.close()

        self.assertEqual(response.status, 200)
        self.assertTrue(content_type.startswith("application/json"))
        self.assertIsNone(content_disposition)
        self.assertEqual(response_body["fileName"], "ego-log_ego_analysis.xlsx")
        export_path = self.export_directory / response_body["fileName"]
        self.assertEqual(response_body["savedPath"], str(export_path))
        self.assertEqual(
            load_workbook(export_path).sheetnames,
            ["运行明细", "违规汇总", "场景违规矩阵"],
        )

    def test_dashboard_uses_json_export_endpoint(self):
        connection = HTTPConnection(
            "127.0.0.1",
            self.server.server_address[1],
            timeout=10,
        )
        connection.request("GET", "/ego/batch")
        response = connection.getresponse()
        dashboard = response.read().decode("utf-8")
        connection.close()

        self.assertEqual(response.status, 200)
        self.assertIn("request(exportUrl, { method: \"POST\" })", dashboard)
        self.assertNotIn("/export.csv", dashboard)
        self.assertNotIn("response.arrayBuffer()", dashboard)


if __name__ == "__main__":
    unittest.main()
